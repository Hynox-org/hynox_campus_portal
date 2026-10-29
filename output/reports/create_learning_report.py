from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path('D:/hynox_campus/output/reports')
stem = 'Hynox-Campus-Modules-7-8-9-Testing-Report-2026-10-29'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(13)
for line in (root / (stem + '.md')).read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:], 'Title')
    elif line.startswith('## '):
        doc.add_heading(line[3:], 1)
    elif line.startswith('- '):
        doc.add_paragraph(line[2:].replace('`', ''), 'List Bullet')
    else:
        doc.add_paragraph(line.replace('`', ''))
sec.footer.paragraphs[0].text = 'Hynox Campus Portal  |  Testing evidence  |  29 October 2026'
sec.footer.paragraphs[0].runs[0].font.size = Pt(9)
doc.save(root / (stem + '.docx'))
print(root / (stem + '.docx'))
