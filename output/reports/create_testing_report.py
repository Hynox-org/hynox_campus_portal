from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path('D:/hynox_campus')
OUT = ROOT / 'output/reports/Hynox-Campus-Modules-4-5-6-Testing-Report-2026-10-27.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    s=doc.styles[name]; s.font.name='Calibri'; s.font.color.rgb=RGBColor(0,0,0)
    for border in s.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(25)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)

def p(text): doc.add_paragraph(text)
def h(text): doc.add_heading(text,2)
def page(title): doc.add_page_break(); doc.add_heading(title,1)
def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for i,v in enumerate(headers): t.rows[0].cells[i].text=v
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for c,v in zip(t.add_row().cells,row): c.text=v
    for ri,row in enumerate(t.rows):
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'243746' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF')); pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3)
                for run in para.runs:
                    run.font.size=Pt(10)
                    if ri==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

doc.add_heading('Hynox Campus Portal Testing Report',0)
p('Modules 4 5 and 6 | 27 October 2026')
p('Prepared for the Hynox Campus Portal owner. This report explains the defects found, the changes made and the evidence collected while testing course blueprints, academic structure and group delivery.')
p('The tested workflows have improved, and the automated suite passes. Production release remains blocked by the deferred database security work and the outstanding verification listed in this report. Passing tests does not establish that every format, device or production condition is safe.')
h('Verification overview')
table(['Check','Result','Meaning'],[
('Automated suite','468 passed in 59 files','Unit, component and service regressions across the repository.'),
('Database integration','Passed targeted fixtures','Synthetic transactions were rolled back; dedicated live QA data was retained separately.'),
('Browser checks','Four dashboard roles exercised','Admin, Institution Admin, Teacher and Student journeys used Brave.'),
('Build and types','See final validation record','Compilation, type checking and lint are separate checks.'),
('Production approval','Blocked','Eleven delivery tables still lack RLS; other open checks remain.')],[1.35,1.75,3.55])
h('What was preserved')
p('Existing fixes and real account roles were preserved. Earlier Modules 1 to 3 were complete according to the supplied handoff, which reported 343 tests in 40 files. Their original Markdown reports were absent from this checkout. The user had already confirmed two tab logout. This report concentrates on the subsequent work in Modules 4 to 6.')
h('How the work was checked')
p('Features were traced from dashboard controls to server actions, services and Supabase. Targeted cases included valid and missing inputs, malformed values, duplicates, campus mismatch, unauthorized callers, deleted parents, stale query data and retries. Failures were reproduced before shared fixes and regression tests were added.')

page('Module 4 Course blueprints and learning assets')
p('Scope covered template editing, publication, ordering, copying a blueprint into a campus and displaying copied lessons and resources to learners.')
table(['Defect found','Change made','Verification'],[
('Copy function accepted an untrusted actor and lacked complete authorization.','Database function checks identity, role, campus and publication; duplicate retries are protected.','Rolled back copy fixtures tested forbidden actors, campus mismatch, drafts, deletion and retry.'),
('Student lesson attachments were missing.','A lesson keyed resource query loads attachments, exposes retry and ignores outdated responses.','Seven links appeared in Student; query and component regressions cover cache and stale results.'),
('Teacher and Student used inconsistent media handling.','Shared renderer handles supported video, audio, images, PDF and safe fallback links.','Teacher playback and Student video/PDF journeys were manually exercised.'),
('PDF pages clipped on narrow screens.','PDF default view fits width while preserving explicit page and zoom settings.','Student 320px PDF retest and Teacher ten page PDF load passed.'),
('Lessons from separate modules interleaved.','Flattened lesson order follows module order then lesson position.','Ordering regression and Teacher eleven lesson navigation passed.'),
('Large blueprint trees fetched too much immediately.','Descendants load on expansion and reuse the query cache.','Request count and cache reuse regression passed.')],[1.8,2.15,2.7])
h('Other corrections')
p('Blueprint runtime validation and role checks were strengthened. Publication controls, missing row errors, URL clearing and explicit zero duration were corrected. Failed OAuth messages and the login button state after browser history restoration were also addressed.')

page('Learning formats and mobile verification')
table(['Format or journey','Observed result','Limit'],[
('Markdown and downloads','Student heading, bullets and seven resource links rendered.','Every document format is not an inline viewer.'),
('YouTube','Teacher and Student playback observed.','Other account restrictions and unavailable videos remain provider dependent.'),
('Uploaded and external MP4','Student playback observed; Teacher uploaded clip reached 5.055 seconds.','Every codec and browser was not tested.'),
('Audio and image','Teacher audio reached one second and image rendered.','Final Student audio/image playback was not repeated; only its new QA cohort was available in the selector.'),
('PDF variants','Portrait, landscape, scanned, fillable and ten page viewers loaded in Student checks.','Form saving, encrypted, corrupt and XFA PDFs were not established.'),
('Teacher responsive preview','Inspected at 320, 375, 430 and 768px; selector overflow fixed.','Brave emulation, not physical iOS or Android devices.'),
('Student responsive lesson','New text lesson rendered at 320px; earlier PDF fit check passed.','Not a claim that every screen size is tested.')],[1.5,2.6,2.55])
h('Storage and upload evidence')
p('A dedicated QA bucket holds the video, audio, image and PDF fixtures. Range requests returned HTTP 206 with range support, CORS and public cache headers. These checks support media delivery but do not measure production throughput or all storage configurations.')
p('Current resource forms accept URLs. QA files were uploaded using a controlled script. A browser file picker upload workflow was not demonstrated, and the report does not describe it as implemented.')
h('Retained course')
p('The Module 4 QA course contains three modules, eleven lessons and seven attachments. Its manifest, upload results and range checks are retained under output/testing/module-04. These are dedicated review fixtures, distinct from rolled back database integration transactions.')

page('Module 5 Academic structure')
p('The hierarchy is program, course, module, lesson and resource. A valid child now derives campus ownership from its saved parent, and invalid or deleted ancestry is rejected. Administrative structure changes and assigned teacher curriculum changes have separate access checks.')
table(['Defect found','Change made','Verification'],[
('Inconsistent hierarchy and scope checks.','Shared academic scope validation plus database integrity triggers.','Mismatch, deleted parent and unauthorized fixtures passed.'),
('Standalone teachers lost institution visibility.','Campus resolver includes explicit, course and cohort assignments.','Cohort only teacher SQL fixtures and live Teacher campus discovery passed.'),
('Teacher prefetch returned too broad a set.','Assigned course and enrolled student filters run before data is returned.','Service/action regressions cover role and assignment scope.'),
('Program publication was difficult to control.','Program visibility and course status/visibility controls added; zero values preserved.','Admin set QA program and course active and campus visible.'),
('Explorer preloaded descendants and could remain stuck after a failed read.','Expansion based loading, cache reuse and loading state recovery.','Lazy loading and partial failure retry regressions passed.'),
('Catalog request errors resembled an empty catalog.','Visible error and Retry control for program/course reads.','New regression failed before the fix and passed afterward.')],[1.8,2.25,2.6])
h('Manual hierarchy journey')
p('Admin created QA Module 5-6 Hierarchy, QA M56 Manual Course, QA M56 Module One and QA M56 Text Lesson through Brave. Required program input was checked. The Student later opened the assigned course and saw the module and Markdown lesson, including its heading and three bullets. The same text lesson was inspected at 320px.')

page('Module 6 Groups and teacher assignments')
p('Groups connect a program, students, courses and teachers. A teacher can remain standalone and join assigned cohorts; additive teacher assignment preserves other teachers. Database guards also protect dates and program relationships.')
table(['Defect found','Change made','Verification'],[
('Campus selection displayed a value while requests used an empty campus ID.','Derived selection stays aligned with late institution props.','Component regression covers late data and valid query scope.'),
('Course assignment picker had no courses.','Load courses for the selected cohort program and clear old selection on cohort changes.','Admin assignment succeeded; loading, retry and stale choice tests passed.'),
('Teacher mapping drawers stayed stale.','Observe query caches and invalidate them on protected revision events.','Open Admin drawer reflected the external QA mapping without refresh.'),
('Delivery queries read broadly then filtered.','Apply campus/cohort predicates before fetching and validate caller roles.','Isolation and caller spoofing regressions passed.'),
('Dates and program changes could invalidate existing delivery.','Reject conflicting date ranges and program changes with dependent records.','Rolled back date, duplicate and enrollment fixtures passed.'),
('Institution analytics counted enrollments as people.','Count unique users and average each person before campus mean; use user ID ranking keys.','Two users across three enrollments yielded two students and 75 percent mean.'),
('Mobile roster selector overflowed.','Stack selection controls on small screens.','320px retest passed; table columns reachable by horizontal scrolling.')],[1.8,2.25,2.6])
p('Admin created QA M56 Collaborative Group, enrolled only the authorized Test Student and assigned its matching program course. Institution Admin saw the campus roster, and Student saw the assigned cohort, course and lesson. Unrelated account enrollments were preserved.')

page('Caching and automatic updates')
h('What changed')
p('Dashboards subscribe to protected revision counters rather than broadcasting student or lesson records. Revision events invalidate the affected query groups. Updates are coalesced for 250 milliseconds, and visible pages check revisions every 30 seconds as a recovery mechanism. Reconnect handling triggers a resynchronization.')
p('Curriculum caches now mirror refreshed query results. Teacher preview and Student selection state follow refreshed data without intentionally replacing an edit in progress. Admin teacher and cohort drawers use observed queries instead of stale local arrays.')
h('Evidence and limits')
table(['Area','Evidence','Practical limit'],[
('Shared editing','Teacher preview reflected an external lesson title change without refresh; original title restored.','Observed within a seven second development check, not a service guarantee.'),
('Teacher mapping','Open Admin mapping drawer picked up an external QA assignment.','Two simultaneous live Teacher browser sessions were not exercised.'),
('Student delivery loading','Thirty cohort regression verifies five batched reads and no progress writes during GET.','Production latency and very large result limits remain unmeasured.'),
('Curriculum trees','Modules, lessons and resources load on expansion with cache reuse.','Large catalog pagination and concurrency still need stress tests.'),
('Media loading','Metadata preload, lazy images/iframes and range delivery checks.','Private URL expiry, all providers and all codecs remain untested.')],[1.45,2.75,2.45])
h('Progress and analytics follow up')
p('Student delivery reads stored course progress rather than recalculating it and writing during every page load. Exact denominator recalculation after curriculum changes needs the learning/progress pass. Institution quiz analytics also need assignment specific eligibility verification; the unique student counting fix does not certify every analytics formula.')

page('Final validation and release conditions')
h('Automated validation')
p('The final automated suite passed 468 tests across 59 files. The suite covers validation, hierarchy ownership, role and campus checks, resource rendering, stale responses, query invalidation, teacher assignment discovery, delivery batching, picker recovery and unique student analytics. Tests are evidence for these cases, not exhaustive production certification.')
p('Build, type and lint results are recorded in the accompanying final validation record below. Existing lint debt is tracked explicitly; the codebase is not described as fully lint clean.')
h('Database integration')
p('Synthetic SQL transactions checked invalid parent/campus relationships, deleted ancestors, duplicate assignments, role eligibility, enrollment isolation, date constraints, atomic instructor replacement, two teachers sharing curriculum and protected revision access. These integration fixtures were rolled back. Dedicated browser QA records remain for inspection.')
h('Open security work')
p('Production release is not approved. Eleven delivery tables lack row level security: cohort_statuses, enrollment_statuses, progress_statuses, lesson_progress, course_assignments, course_progress, enrollments, cohorts, attendance_sessions, cohort_teachers and attendance_records. The user deferred this broader RLS pass. Application checks are not a replacement for database policies.')
p('Also review permissive academic read policies, remaining database definer functions and the security/performance advisor findings. The saved advisor summary includes unindexed foreign keys and policy warnings; these require workload and access review rather than blanket changes.')
h('Untested production conditions')
p('Remaining work includes physical phones and Safari, corrupt or encrypted PDFs, PDF form saving, private media expiry, full offline recovery, concurrent edit conflicts, large catalogs and PostgREST limits, realtime quotas, deployment limits and production load measurements. Original earlier deployment evidence was unavailable in the requested Markdown paths.')
h('Evidence locations')
p('Module 4: output/testing/module-04/current-verification.md and live-qa-manifest.json. Module 5: output/testing/module-05/report.md, database-integration.sql, advisor-summary.json and final validation logs. Module 6: output/testing/module-06/report.md and assignment/analytics regression logs. The consolidated scope ledger is output/testing/dashboard-scope-and-remaining.md.')

# Actual final command outcomes are appended by the verification step before rendering.
validation=ROOT/'output/testing/module-05/final-validation-summary.txt'
if validation.exists():
    h('Final command results')
    p(validation.read_text(encoding='utf-8').strip())
doc.core_properties.title='Hynox Campus Portal Testing Report'
doc.core_properties.subject='Modules 4 5 and 6 verification and fixes'
doc.core_properties.author='Hynox Campus QA'
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
